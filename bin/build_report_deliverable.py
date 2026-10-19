#!/usr/bin/env python3
"""Build the "WS1 Promo Codes by Pillar" deliverable workbook (the May/Jascinta sheet) from a report
payload + the seeded definitions. One self-contained builder — replaces the 3 scratchpad scripts.

Usage: python bin/build_report_deliverable.py --payload <report.json> --defs <report-definitions.json> --out <xlsx>
Stamps the payload's window as the "as of" basis so freshness is always visible in the sheet.
"""
import json, sys, argparse, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ap = argparse.ArgumentParser()
ap.add_argument("--payload", required=True)
ap.add_argument("--defs", required=True)
ap.add_argument("--out", required=True)
A = ap.parse_args()

d = json.load(open(A.payload, encoding="utf-8"))
DEFS = json.load(open(A.defs, encoding="utf-8"))
PERIOD = d.get("period", "")          # e.g. "2026-01-01 .. 2026-10-01" (END_EXCL)
# derive a human window label: END_EXCL -> inclusive end
def human_window(period):
    try:
        a, b = [x.strip() for x in period.split("..")]
        from datetime import date, timedelta
        end_incl = date.fromisoformat(b) - timedelta(days=1)
        return f"{a} to {end_incl.isoformat()}"
    except Exception:
        return period
WINDOW = human_window(PERIOD)
BUILT = d.get("builtAt", "")

NAVY="22304A"; HEAD_FG="FFFFFF"; BAND="F4F6F9"; SECT="2F5D50"; MUT="52606D"
PILL_FILL={"Acquisition":"E8F0FE","Retention":"E6F4EA","VIP":"FCE8E6"}
thin=Side(style="thin",color="D5DBE2"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
PILLAR_ORDER={"Acquisition":0,"Retention":1,"VIP":2}
COLS=["Pillar","Promo code","Bonus type","Call (verdict)","Bonus spend","Players","Bonus ROI · 7d","Value band","Tier"]
WIDTHS=[13,32,16,15,15,10,13,12,26]

def code_rows(mk):
    m=d.get(mk) or {}; sym=m.get("sym",""); out=[]
    for c in (m.get("codes") or []):
        code=(c.get("cells",{}).get("Promo code") or {}).get("d") or "?"
        out.append({"Pillar":c.get("_pillar","?"),"Promo code":code,"Bonus type":c.get("_type",""),
                    "Call (verdict)":c.get("_call",""),"Bonus spend":c.get("_spend"),"Players":c.get("_players"),
                    "Bonus ROI · 7d":c.get("_perrm"),"Value band":c.get("_band",""),"Tier":c.get("_tier","")})
    out.sort(key=lambda r:(PILLAR_ORDER.get(r["Pillar"],9),-(r["Bonus spend"] or 0)))
    return out, sym

wb=Workbook(); wb.remove(wb.active)

def market_tab(mk):
    rows,sym=code_rows(mk); ws=wb.create_sheet(f"{mk} — codes by pillar")
    ws.cell(1,1,f"{mk} promo codes by pillar").font=Font(bold=True,size=13,color=NAVY)
    ws.cell(2,1,f"Year-to-date · {WINDOW} · {len(rows)} graded codes · spend in {sym}").font=Font(italic=True,size=10,color=MUT)
    hr=4
    for j,name in enumerate(COLS,1):
        cc=ws.cell(hr,j,name); cc.fill=PatternFill("solid",fgColor=NAVY); cc.font=Font(bold=True,color=HEAD_FG,size=10)
        cc.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cc.border=BORDER
    for i,r in enumerate(rows):
        rr=hr+1+i
        for j,name in enumerate(COLS,1):
            v=r[name]; cell=ws.cell(rr,j)
            if name=="Bonus spend" and isinstance(v,(int,float)): cell.value=v; cell.number_format=f'"{sym}"#,##0'
            elif name=="Players" and isinstance(v,(int,float)): cell.value=v; cell.number_format="#,##0"
            elif name=="Bonus ROI · 7d": cell.value=(f"{v:+.2f}" if isinstance(v,(int,float)) else "—")
            else: cell.value=("" if v is None else v)
            cell.border=BORDER; cell.font=Font(size=10)
            if name in ("Bonus spend","Players","Bonus ROI · 7d"): cell.alignment=Alignment(horizontal="right")
        ws.cell(rr,1).fill=PatternFill("solid",fgColor=PILL_FILL.get(r["Pillar"],"FFFFFF"))
        if i%2==1:
            for j in range(2,len(COLS)+1):
                if ws.cell(rr,j).fill.fgColor.rgb in (None,"00000000"): ws.cell(rr,j).fill=PatternFill("solid",fgColor=BAND)
    for j,w in enumerate(WIDTHS,1): ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes=f"A{hr+1}"; ws.auto_filter.ref=f"A{hr}:{get_column_letter(len(COLS))}{hr+len(rows)}"
    return rows

def summary(my_rows,sg_rows):
    ws=wb.create_sheet("Read me + Summary",0)
    ws.cell(1,1,"WS1 promo codes by pillar — year-to-date").font=Font(bold=True,size=14,color=NAVY)
    ws.cell(2,1,f"AS OF: {WINDOW}  ·  report built {BUILT[:10]}.  These are the GRADED (managed) codes the report scores with a firm call.").font=Font(italic=True,size=10,color=MUT)
    hdr=["Pillar","MY codes","MY bonus spend (RM)","SG codes","SG bonus spend (S$)"]; hr=4
    for j,name in enumerate(hdr,1):
        cc=ws.cell(hr,j,name); cc.fill=PatternFill("solid",fgColor=NAVY); cc.font=Font(bold=True,color=HEAD_FG,size=10)
        cc.alignment=Alignment(horizontal="center",wrap_text=True); cc.border=BORDER
    agg=lambda rows,p:(len([r for r in rows if r["Pillar"]==p]),sum((r["Bonus spend"] or 0) for r in rows if r["Pillar"]==p))
    for i,p in enumerate(["Acquisition","Retention","VIP"]):
        rr=hr+1+i; mn,msp=agg(my_rows,p); sn,ssp=agg(sg_rows,p)
        ws.cell(rr,1,p).fill=PatternFill("solid",fgColor=PILL_FILL.get(p,"FFFFFF"))
        ws.cell(rr,2,mn).number_format="#,##0"; ws.cell(rr,3,msp).number_format='"RM"#,##0'
        ws.cell(rr,4,sn).number_format="#,##0"; ws.cell(rr,5,ssp).number_format='"S$"#,##0'
        for j in range(1,6): ws.cell(rr,j).border=BORDER
    tr=hr+4
    ws.cell(tr,1,"Total").font=Font(bold=True)
    ws.cell(tr,2,len(my_rows)).number_format="#,##0"; ws.cell(tr,3,sum((r["Bonus spend"] or 0) for r in my_rows)).number_format='"RM"#,##0'
    ws.cell(tr,4,len(sg_rows)).number_format="#,##0"; ws.cell(tr,5,sum((r["Bonus spend"] or 0) for r in sg_rows)).number_format='"S$"#,##0'
    for j in range(1,6): ws.cell(tr,j).font=Font(bold=True); ws.cell(tr,j).border=BORDER
    for j,w in enumerate([13,11,20,11,20],1): ws.column_dimensions[get_column_letter(j)].width=w
    note=["See the Definitions tab for the glossary, and the Validation reference tab to reconcile every metric against the warehouse.",
          "GRADED = codes with a firm call (Scale/Maintain/Optimise/Reduce/Stop). More codes run than appear here; the rest are out of the ROI-graded scope (VIP perks, reload, weekly-rescue, referral, low-volume, ad-hoc) — this is a scope split, not a maturity one.",
          f"This sheet auto-refreshes monthly from the latest report build. AS OF date above shows how current it is."]
    for k,t in enumerate(note): ws.cell(hr+7+k,1,"• "+t).font=Font(size=9,color=MUT)
    ws.column_dimensions["A"].width=118 if False else ws.column_dimensions["A"].width

def defs_tab():
    """Definitions glossary — formulas + callsLevers come from the payload (stable); the rest is fixed."""
    my=d.get("MY") or {}
    ws=wb.create_sheet("Definitions",1)
    for j,w in enumerate([30,46,40,40],1): ws.column_dimensions[get_column_letter(j)].width=w
    r=[1]
    def row(cells,bold=False,size=10,fill=None,fg=None,border=False,italic=False):
        for j in range(1,5):
            v=cells[j-1] if j-1<len(cells) else ""
            c=ws.cell(r[0],j,v); c.font=Font(bold=bold,size=size,color=(fg or "1F2933"),italic=italic)
            c.alignment=Alignment(wrap_text=True,vertical="top")
            if fill: c.fill=PatternFill("solid",fgColor=fill)
            if border: c.border=BORDER
        r[0]+=1
    def section(t):
        r[0]+=1; c=ws.cell(r[0],1,t); c.font=Font(bold=True,size=12,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=SECT)
        for j in range(2,5): ws.cell(r[0],j).fill=PatternFill("solid",fgColor=SECT)
        r[0]+=1
    ws.cell(r[0],1,"WS1 Promo Effectiveness — definitions").font=Font(bold=True,size=14,color=NAVY); r[0]+=1
    ws.cell(r[0],1,f"Every term as the dashboard uses it. AS OF {WINDOW}. Metrics/verdicts/formulas pulled from the report.").font=Font(italic=True,size=10,color=MUT); r[0]+=2
    section("1 · How a code gets onto the dashboard")
    row(["Term","Definition"],bold=True,fill=NAVY,fg=HEAD_FG,border=True)
    for t,x in [("Graded / managed code","A code the report scores with a firm call (Scale/Maintain/Optimise/Reduce/Stop). Only graded codes appear on Summary & Codes."),
                ("Mature (claim)","A claim with enough follow-up BEFORE the window end to measure its effect: 7d min (any call), 30d (firm call + all VIP), ~90d (at-risk)."),
                ("Window-end rule","Maturity is to the window END (AS_OF), not today. mature_before(N)=END_EXCL−(N−1)."),
                ("Scope (left out)","Not ROI-graded: VIP perks/birthday FC, weekly-rescue/win-back, reload, referral, low-volume, ad-hoc. They still show in BO.")]:
        row([t,x],border=True)
    section("2 · Metrics & economics (from the report Reference tab)")
    fm=my.get("formulas") or []
    for rr in fm:
        if not isinstance(rr,list): continue
        h=str(rr[0]).strip()
        if h=="Metric": row(["Metric","Formula","Worked example","How to read it"],bold=True,fill=NAVY,fg=HEAD_FG,border=True)
        elif h and h[0] in "①②③④": row([h,"","",""],bold=True,fill="EAF0F2")
        elif h: row([rr[0],rr[1] if len(rr)>1 else "",rr[2] if len(rr)>2 else "",rr[3] if len(rr)>3 else ""],border=True,size=9)
    section("3 · The call / verdict")
    cl=my.get("callsLevers") or []; started=False
    for rr in cl:
        if not isinstance(rr,list): continue
        h=str(rr[0]).strip()
        if h in ("The call","The six levers","Two things to remember"):
            if h=="The call": row(["The call","What it means","Action type","Specific lever(s)"],bold=True,fill=NAVY,fg=HEAD_FG,border=True)
            else: row([h,"","",""],bold=True,fill="EAF0F2")
            started=True; continue
        if started and h: row([rr[0],rr[1] if len(rr)>1 else "",rr[2] if len(rr)>2 else "",rr[3] if len(rr)>3 else ""],border=True)
    section("4 · Value bands & recency states")
    row(["Term","Definition"],bold=True,fill=NAVY,fg=HEAD_FG,border=True)
    for t,x in [("Value band — at-risk view","By YTD deposits: Top value ≥100k · High ≥20k · Mid ≥2k · Low. (SG uses same numbers in S$.)"),
                ("Value band — config view","By last-90-day deposits: Whale ≥50k · High ≥5k · Mid ≥500. Not comparable to the at-risk view."),
                ("Recency state","Days since last deposit (to window end): Active ≤30 · Cooling 31–60 · Lapsed 61–90 · Dormant >90.")]:
        row([t,x],border=True)

def validation_tab():
    def title_for(a):
        s=a.lower()
        if "verdict decision tree" in s: return (0,"A · Verdict decision tree (per pillar)")
        if "attribution" in s: return (1,"B · Attribution internals (NGR lift / Bonus ROI)")
        if "deposits back" in s or "deposit timing" in s: return (2,"C · Deposits-back family & deposit timing")
        if "keeps players" in s: return (3,"D · Keeps players, Review & Money-tab columns")
        if "acquisition" in s: return (4,"E · Acquisition metrics")
        if "codes enter" in s or "universe" in s: return (5,"F · Code universe & classification")
        if "vip driver" in s: return (6,"G · VIP lanes & driver metrics (all 4 lanes)")
        if "derived text" in s: return (7,"H · Derived text columns & cross-references")
        return (9,a[:60])
    ordered=sorted(DEFS,key=lambda a:title_for(a.get("area",""))[0])
    ws=wb.create_sheet("Validation reference",2)
    HDR=["Metric","Exact definition","Source & filters","Window / maturity","Caveat"]
    for j,w in enumerate([26,62,46,30,40],1): ws.column_dimensions[get_column_letter(j)].width=w
    r=[1]
    def put(cells,bold=False,size=9,fill=None,fg=None,border=False,italic=False):
        for j in range(1,6):
            v=cells[j-1] if j-1<len(cells) else ""
            c=ws.cell(r[0],j,v); c.font=Font(bold=bold,size=size,color=(fg or "1F2933"),italic=italic)
            c.alignment=Alignment(wrap_text=True,vertical="top")
            if fill: c.fill=PatternFill("solid",fgColor=fill)
            if border: c.border=BORDER
        r[0]+=1
    ws.cell(r[0],1,"WS1 Promo Effectiveness — validation reference").font=Font(bold=True,size=14,color=NAVY); r[0]+=1
    put(["Every metric with exact formula, warehouse source + filters, window/maturity and caveat — reproduce each figure. Example dates use the pipeline default window; for this sheet AS_OF = the window end.","","","",""],italic=True,fg=MUT)
    r[0]+=1
    for a in ordered:
        _,title=title_for(a.get("area","")); c=ws.cell(r[0],1,title); c.font=Font(bold=True,size=12,color="FFFFFF"); c.fill=PatternFill("solid",fgColor=SECT)
        for j in range(2,6): ws.cell(r[0],j).fill=PatternFill("solid",fgColor=SECT)
        r[0]+=1
        put(HDR,bold=True,fill=NAVY,fg=HEAD_FG,border=True)
        for de in (a.get("definitions") or []):
            put([de.get("column",""),de.get("definition",""),de.get("source",""),de.get("window_maturity",""),de.get("caveat","")],border=True)
        r[0]+=1

my_rows=market_tab("MY"); sg_rows=market_tab("SG")
summary(my_rows,sg_rows); defs_tab(); validation_tab()
# tab order: Read me, Definitions, Validation reference, MY, SG
wb._sheets.sort(key=lambda s:{"Read me + Summary":0,"Definitions":1,"Validation reference":2,"MY — codes by pillar":3,"SG — codes by pillar":4}.get(s.title,9))
wb.save(A.out)
sys.stdout=io.TextIOWrapper(sys.stdout.buffer,encoding="utf-8")
print(json.dumps({"ok":True,"out":A.out,"window":WINDOW,"period":PERIOD,"builtAt":BUILT,"my_codes":len(my_rows),"sg_codes":len(sg_rows)}))
